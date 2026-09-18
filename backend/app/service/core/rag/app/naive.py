#
#  Copyright 2025 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#

import logging
import os
from dataclasses import dataclass
from tika import parser
from io import BytesIO
from docx import Document
from timeit import default_timer as timer
import re
from html import escape
from service.core.deepdoc.parser.pdf_parser import PlainParser
from service.core.rag.nlp import rag_tokenizer, naive_merge, tokenize_table, tokenize_chunks, find_codec, concat_img, \
    naive_merge_docx, tokenize_chunks_docx
from service.core.deepdoc.parser import PdfParser, ExcelParser, DocxParser, HtmlParser, JsonParser, MarkdownParser, TxtParser
from service.core.rag.utils import num_tokens_from_string
from PIL import Image
from functools import reduce
from markdown import markdown
from docx.image.exceptions import UnrecognizedImageError, UnexpectedEndOfFileError, InvalidImageStreamError


@dataclass
class _DocxParseState:
    """保存段落解析过程中需要传给下一段的临时状态。"""

    lines: list
    pending_images: list
    pending_caption_index: int = None


class Docx(DocxParser):
    def __init__(self):
        pass

    def get_picture(self, document, paragraph):
        """提取段落中的第一张图片，无法识别时返回 None。"""
        # 先从段落底层 XML 中寻找图片节点。普通文本段落不会有 pic:pic。
        img = paragraph._element.xpath('.//pic:pic')
        if not img:
            return None

        # 当前流程每个段落只处理第一张图片。
        img = img[0]

        # r:embed 是图片在 DOCX 关系表中的编号，通过它找到真正的图片数据。
        embed = img.xpath('.//a:blip/@r:embed')[0]
        related_part = document.part.related_parts[embed]
        try:
            image_blob = related_part.image.blob
        except UnrecognizedImageError:
            logging.info("Unrecognized image format. Skipping image.")
            return None
        except UnexpectedEndOfFileError:
            logging.info("EOF was unexpectedly encountered while reading an image stream. Skipping image.")
            return None
        except InvalidImageStreamError:
            logging.info("The recognized image stream appears to be corrupted. Skipping image.")
            return None
        try:
            # 统一转成 RGB，方便后面把多张图片纵向拼接。
            image = Image.open(BytesIO(image_blob)).convert('RGB')
            return image
        except Exception:
            # Pillow 无法打开的图片不会中断整份文档解析，直接跳过。
            return None

    def __clean(self, line):
        """统一段落空格并删除首尾空白。"""
        line = re.sub(r"\u3000", " ", line).strip()
        return line

    def _open_document(self, filename, binary):
        """优先解析调用方传入的二进制内容，否则从文件路径读取 DOCX。"""
        source = filename if not binary else BytesIO(binary)
        return Document(source)

    @staticmethod
    def _style_name(paragraph):
        """安全地获取段落样式名称；没有样式时返回空字符串。"""
        return paragraph.style.name if paragraph.style else ""

    @staticmethod
    def _is_caption(paragraph):
        """判断当前段落是不是 Word 的图片题注。"""
        return Docx._style_name(paragraph) == "Caption"

    @staticmethod
    def _take_previous_inline_image(lines):
        """取出上一段普通正文自身携带的最后一张图片。"""
        if not lines or lines[-1][2] == "Caption":
            return None

        images = lines[-1][1]
        if images and images[-1] is not None:
            return images.pop()
        return None

    def _append_caption(self, paragraph, state):
        """
        保存题注，并优先匹配题注前面尚未确定归属的图片。

        如果前面没有图片，就记住当前题注在 lines 中的位置，等待后面的
        纯图片段落，从而同时支持“图片 -> 题注”和“题注 -> 图片”。
        """
        caption_images = state.pending_images
        state.pending_images = []

        # 图片也可能与上一段文字同处一个 Word 段落。这类图片不会进入
        # pending_images，因此在没有纯图片时，再从上一段取一张图片。
        if not caption_images:
            inline_image = self._take_previous_inline_image(state.lines)
            if inline_image is not None:
                caption_images = [inline_image]

        state.lines.append((
            self.__clean(paragraph.text),
            caption_images or [None],
            self._style_name(paragraph),
        ))

        # 没有匹配到前置图片时，让当前题注等待后面的图片。
        state.pending_caption_index = (
            None if caption_images else len(state.lines) - 1
        )

    @staticmethod
    def _attach_unmatched_images(state, target_images):
        """把没等到题注的图片交给兜底段落，并清空待匹配图片。"""
        target_images.extend(state.pending_images)
        state.pending_images = []

    def _append_text_paragraph(self, paragraph, state):
        """保存普通文本段落，并结束此前尚未完成的题注匹配。"""
        # get_picture 当前只提取段落里的第一张图片；没有图片时返回 None。
        images = [self.get_picture(self.doc, paragraph)]

        if state.pending_images:
            if state.lines:
                # 图片后出现的是普通正文而不是题注：沿用原来的兜底规则，
                # 把图片附加到它前面的最近一段文字。
                self._attach_unmatched_images(state, state.lines[-1][1])
            else:
                # 文档开头先出现图片时，把图片交给随后出现的第一段正文。
                leading_images = state.pending_images
                state.pending_images = []
                images = leading_images + images

        # 普通正文会中断“题注等待后续图片”的相邻关系，避免后面的图片
        # 跨过正文错误匹配到更早的题注。
        state.pending_caption_index = None

        # lines 的中间结构为：(段落文本, 图片列表, 段落样式名)。
        state.lines.append((
            self.__clean(paragraph.text),
            images,
            self._style_name(paragraph),
        ))

    def _append_image_paragraph(self, paragraph, state):
        """处理纯图片段落，并尝试匹配图片前面的待处理题注。"""
        current_image = self.get_picture(self.doc, paragraph)
        if not current_image:
            return

        if state.pending_caption_index is not None:
            # “题注 -> 图片”：当前图片直接交给前面等待匹配的题注。这里
            # 不立即清除题注索引，因此连续出现的多张图片会归于同一题注；
            # 遇到普通正文或下一个题注时，等待状态才会结束。
            caption = state.lines[state.pending_caption_index]
            if caption[1] == [None]:
                caption[1].clear()
            caption[1].append(current_image)
            return

        # “图片 -> 题注”：先暂存图片。下一条有效文本若为题注，题注会
        # 取走图片；若为普通正文，则使用普通段落的兜底规则。
        state.pending_images.append(current_image)

    def _append_paragraph(self, paragraph, state):
        """根据段落内容和样式，把段落交给对应的处理方法。"""
        # 没有文字的段落可能是纯图片段落，也可能是真正的空段落。
        if not paragraph.text.strip():
            self._append_image_paragraph(paragraph, state)
            return

        # Caption 表示图片或表格附近的题注文字。
        if self._is_caption(paragraph):
            self._append_caption(paragraph, state)
            return

        self._append_text_paragraph(paragraph, state)

    @staticmethod
    def _page_break_count(paragraph):
        """统计一个段落中可以识别到的分页标记数量。"""
        count = 0
        for run in paragraph.runs:
            xml = run._element.xml

            # Word 保存文件时记录的“上次渲染分页位置”。
            if "lastRenderedPageBreak" in xml:
                count += 1

            # 用户在 Word 中主动插入的分页符。
            elif "w:br" in xml and 'type="page"' in xml:
                count += 1
        return count

    def _extract_lines(self, from_page, to_page):
        """
        按文档顺序提取指定页码范围内的段落和图片。

        page_number 从 0 开始，解析范围是 [from_page, to_page)。DOCX 的
        自动分页依赖 Word 排版，因此这里只能根据 XML 中已有的分页标记计数。
        """
        page_number = 0
        state = _DocxParseState(lines=[], pending_images=[])

        for paragraph in self.doc.paragraphs:
            if page_number > to_page:
                break
            if from_page <= page_number < to_page:
                self._append_paragraph(paragraph, state)

            # 分页符可能出现在不需要提取的段落中，但仍需统计，才能正确
            # 判断后续段落是否进入目标页码范围。
            page_number += self._page_break_count(paragraph)

        # 文档结束时仍没等到题注的图片，沿用旧逻辑附加到最近一段文字。
        if state.pending_images and state.lines:
            self._attach_unmatched_images(state, state.lines[-1][1])

        # 一个 chunk 最终只能携带一张图片，因此把同段落关联的多张图片
        # 纵向拼接；同时丢弃仅供解析过程使用的段落样式。
        return [
            (text, reduce(concat_img, images) if images else None)
            for text, images, _style in state.lines
        ]

    @staticmethod
    def _row_to_html(row):
        """把一行 Word 表格转换成 HTML 的 ``<tr>``。"""
        html = "<tr>"
        cell_index = 0
        while cell_index < len(row.cells):
            span = 1
            cell = row.cells[cell_index]

            # python-docx 读取横向合并单元格时，可能返回多个内容相同的
            # cell。这里沿用原来的策略，将它们表示成一个 colspan 单元格。
            for following_index in range(cell_index + 1, len(row.cells)):
                if cell._tc is row.cells[following_index]._tc:
                    span += 1
                    cell_index = following_index
                else:
                    break
            cell_index += 1
            html += (
                f"<td>{escape(cell.text)}</td>"
                if span == 1
                else f"<td colspan='{span}'>{escape(cell.text)}</td>"
            )
        return html + "</tr>"

    def _extract_tables(self):
        """
        将文档中的每张表格转换成独立的 HTML 表格块。

        表格和正文分别返回，后续流程会为表格单独生成检索 chunk。
        """
        tables = []
        for table in self.doc.tables:
            html = "<table>" + "".join(
                self._row_to_html(row) for row in table.rows
            ) + "</table>"
            tables.append(((None, html), ""))
        return tables

    def __call__(self, filename, binary=None, from_page=0, to_page=100000):
        """
        解析 DOCX，返回正文段落和表格。

        正文元素的结构是 ``(文本, 图片或 None)``；表格会转换成 HTML，
        放在单独的列表中返回。调用方随后会分别对正文和表格进行切块、分词。
        """
        # 第一步：从文件路径或内存二进制打开 DOCX。
        self.doc = self._open_document(filename, binary)

        # 第二步：提取正文/图片；第三步：提取表格。两者保持原有独立结构。
        return self._extract_lines(from_page, to_page), self._extract_tables()


class Pdf(PdfParser):
    def __call__(self, filename, binary=None, from_page=0,
                 to_page=100000, zoomin=3, callback=None):
        """
        深度解析 PDF，返回普通文本段落和单独提取出的表格/图片。

        这里的 __call__ 是 Python 的特殊方法，因此可以写成 Pdf()(文件名)，
        实际上执行的就是 Pdf().__call__(文件名)。

        :param filename: PDF 文件路径
        :param binary: 可选的 PDF 二进制内容；有值时不再从 filename 读文件
        :param from_page: 从哪一页开始解析，页码从 0 开始
        :param to_page: 解析到哪一页，不包含这一页
        :param zoomin: PDF 转图片时的放大倍数；越大越清晰，但速度和内存开销越高
        :param callback: 向调用者报告解析进度的函数
        :return: (sections, tables)，也就是“普通文本段落列表”和“表格/图片列表”
        """
        # 统计当前步骤耗时；first_start 用来统计完整流程耗时。
        start = timer()
        first_start = start

        # 第一步：PDF 页面转图片并进行 OCR。
        # 如果 PDF 原本就有文字，会同时利用 PDF 自带文字；扫描件则主要依靠 OCR。
        # 结果暂存在 self.boxes 中，每个 box 表示一小块识别出的文字及其坐标。
        callback(msg="OCR started")
        self.__images__(
            # binary 有内容时直接解析内存数据，否则根据 filename 读取磁盘文件。
            filename if not binary else binary,
            zoomin,
            from_page,
            to_page,
            callback
        )
        callback(msg="OCR finished ({:.2f}s)".format(timer() - start))
        logging.info("OCR({}~{}): {:.2f}s".format(from_page, to_page, timer() - start))

        # 第二步：版面分析。
        # 判断各文字框属于标题、正文、页眉、页脚、表格、图片说明等哪种区域。
        # 这里主要是在现有 self.boxes 上补充版面类型，不是重新读取 PDF。
        start = timer()
        self._layouts_rec(zoomin)
        # 0.63 是项目预估的总体进度 63%，不是本步骤完成了 63%。
        callback(0.63, "Layout analysis ({:.2f}s)".format(timer() - start))

        # 第三步：表格结构分析。
        # 识别表格中的行、列、单元格以及单元格之间的关系。
        start = timer()
        self._table_transformer_job(zoomin)
        callback(0.65, "Table analysis ({:.2f}s)".format(timer() - start))

        # 第四步：合并零散文字。
        # OCR 往往按单词或小文本框返回结果，这里把属于同一行、同一段的内容合起来。
        start = timer()
        self._text_merge()
        callback(0.67, "Text merged ({:.2f}s)".format(timer() - start))

        # 第五步：把表格和图片从普通正文中单独提取出来。
        # tbls 后续会由 tokenize_table() 处理，不和普通正文混成一个 chunk。
        tbls = self._extract_table_figure(True, zoomin, True, True)

        # 第六步：按照阅读顺序向下连接相邻文本框，形成更完整的段落。
        # self._naive_vertical_merge()
        self._concat_downward()
        # self._filter_forpages()

        logging.info("layouts cost: {}s".format(timer() - first_start))

        # self.boxes 是最终保留下来的正文文本框。
        # _line_tag() 会给每段文字附加页码及坐标，例如：
        return [(b["text"], self._line_tag(b, zoomin))
                for b in self.boxes], tbls


class Markdown(MarkdownParser):
    def __call__(self, filename, binary=None):
        if binary:
            encoding = find_codec(binary)
            txt = binary.decode(encoding, errors="ignore")
        else:
            with open(filename, "r") as f:
                txt = f.read()
        remainder, tables = self.extract_tables_and_remainder(f'{txt}\n')

        sections = []
        tbls = []
        for sec in remainder.split("\n"):
            if num_tokens_from_string(sec) > 3 * self.chunk_token_num:
                sections.append((sec[:int(len(sec) / 2)], ""))
                sections.append((sec[int(len(sec) / 2):], ""))
            else:
                if sec.strip().find("#") == 0:
                    sections.append((sec, ""))
                elif sections and sections[-1][0].strip().find("#") == 0:
                    sec_, _ = sections.pop(-1)
                    sections.append((sec_ + "\n" + sec, ""))
                else:
                    sections.append((sec, ""))

        for table in tables:
            tbls.append(((None, markdown(table, extensions=['markdown.extensions.tables'])), ""))
        return sections, tbls


DEFAULT_PARSER_CONFIG = {
    "chunk_token_num": 128,
    "delimiter": "\n!?。；！？",
    "layout_recognize": "DeepDOC",
}
DEFAULT_TEXT_DELIMITER = "\n!?;。；！？"
TEXT_FILE_EXTENSIONS = {
    ".txt", ".py", ".js", ".java", ".c", ".cpp", ".h", ".php",
    ".go", ".ts", ".sh", ".cs", ".kt", ".sql",
}


@dataclass
class ParsedDocument:
    """
    保存“文件解析阶段”的统一结果。

    不同解析器返回的数据并不完全相同：PDF 有位置解析器，DOCX 有图片，
    Markdown 和 PDF 可能有表格。把差异放在这个对象中，后面的切块代码
    就不需要再判断每种文件格式。
    """

    sections: list
    tables: list
    pdf_parser: object = None
    uses_docx_merge: bool = False


def _empty_callback(prog=None, msg=""):
    """调用者没有提供进度回调时使用，什么也不做。"""
    pass


def _get_parser_config(kwargs):
    """取得调用者传入的解析配置；没传时使用默认配置。"""
    parser_config = kwargs.get("parser_config")
    if parser_config is None:
        return DEFAULT_PARSER_CONFIG.copy()
    return parser_config


def _build_document_metadata(filename):
    """创建每个最终 chunk 都会复制一份的文档名称和标题分词。"""
    title_tokens = rag_tokenizer.tokenize(
        re.sub(r"\.[a-zA-Z]+$", "", filename)
    )
    return {
        "docnm_kwd": filename,
        "title_tks": title_tokens,
        #更小粒度的标题分词，供检索时使用
        "title_sm_tks": rag_tokenizer.fine_grained_tokenize(title_tokens),
    }


def _load_binary(filename, binary):
    """优先使用已有的二进制内容；没有时再从文件路径读取。"""
    if binary is not None:
        return binary

    with open(filename, "rb") as source_file:
        return source_file.read()


def _parse_docx(filename, binary, callback):
    """解析 DOCX，同时保留段落中提取出来的图片和表格。"""
    callback(0.1, "Start to parse.")
    sections, tables = Docx()(filename, binary)
    callback(0.8, "Finish parsing.")
    return ParsedDocument(
        sections=sections,
        tables=tables,
        uses_docx_merge=True,
    )


def _parse_pdf(filename, binary, from_page, to_page, parser_config, callback):
    """根据配置选择 DeepDOC 或纯文本 PDF 解析器。"""
    pdf_parser = Pdf()
    if parser_config.get("layout_recognize", "DeepDOC") == "Plain Text":
        pdf_parser = PlainParser()

    # binary 不为 None 时，说明文件内容已经在内存中，不需要再次读磁盘。
    source = binary if binary is not None else filename
    sections, tables = pdf_parser(
        source,
        from_page=from_page,
        to_page=to_page,
        callback=callback,
    )
    return ParsedDocument(
        sections=sections,
        tables=tables,
        pdf_parser=pdf_parser,
    )


def _parse_excel(filename, binary, parser_config, callback):
    """解析 XLS/XLSX；html4excel=True 时把每块表格保存成 HTML。"""
    callback(0.1, "Start to parse.")
    file_content = _load_binary(filename, binary)
    excel_parser = ExcelParser()

    if parser_config.get("html4excel"):
        raw_sections = excel_parser.html(file_content, 12)
    else:
        raw_sections = excel_parser(file_content)

    sections = [(section, "") for section in raw_sections if section]
    return ParsedDocument(sections=sections, tables=[])


def _parse_text(filename, binary, parser_config, callback):
    """解析普通文本和代码文件。"""
    callback(0.1, "Start to parse.")
    sections = TxtParser()(
        filename,
        binary,
        parser_config.get("chunk_token_num", 128),
        parser_config.get("delimiter", DEFAULT_TEXT_DELIMITER),
    )
    callback(0.8, "Finish parsing.")
    return ParsedDocument(sections=sections, tables=[])


def _parse_markdown(filename, binary, parser_config, callback):
    """解析 Markdown，并把 Markdown 表格单独提取出来。"""
    callback(0.1, "Start to parse.")
    chunk_token_num = int(parser_config.get("chunk_token_num", 128))
    sections, tables = Markdown(chunk_token_num)(filename, binary)
    callback(0.8, "Finish parsing.")
    return ParsedDocument(sections=sections, tables=tables)


def _parse_html(filename, binary, callback):
    """从 HTML 中提取网页标题和主要正文。"""
    callback(0.1, "Start to parse.")
    raw_sections = HtmlParser()(filename, binary)
    sections = [(section, "") for section in raw_sections if section]
    callback(0.8, "Finish parsing.")
    return ParsedDocument(sections=sections, tables=[])


def _parse_json(filename, binary, parser_config, callback):
    """读取 JSON，并按对象结构拆成多个文本段落。"""
    callback(0.1, "Start to parse.")
    file_content = _load_binary(filename, binary)
    chunk_token_num = int(parser_config.get("chunk_token_num", 128))
    raw_sections = JsonParser(chunk_token_num)(file_content)
    sections = [(section, "") for section in raw_sections if section]
    callback(0.8, "Finish parsing.")
    return ParsedDocument(sections=sections, tables=[])


def _parse_legacy_doc(filename, binary, callback):
    """使用 Apache Tika 解析旧版 .doc 文件。"""
    callback(0.1, "Start to parse.")
    file_content = _load_binary(filename, binary)
    parsed_result = parser.from_buffer(BytesIO(file_content))
    content = parsed_result.get("content")

    if content is None:
        message = f"tika.parser got empty content from {filename}."
        callback(0.8, message)
        logging.warning(message)
        return ParsedDocument(sections=[], tables=[])

    sections = [(section, "") for section in content.split("\n") if section]
    callback(0.8, "Finish parsing.")
    return ParsedDocument(sections=sections, tables=[])


def _parse_document(
    filename,
    binary,
    from_page,
    to_page,
    parser_config,
    callback,
):
    """根据文件扩展名，把任务交给对应的格式解析函数。"""
    file_extension = os.path.splitext(filename)[1].lower()

    if file_extension == ".docx":
        return _parse_docx(filename, binary, callback)
    if file_extension == ".pdf":
        return _parse_pdf(
            filename,
            binary,
            from_page,
            to_page,
            parser_config,
            callback,
        )
    if file_extension in {".xls", ".xlsx"}:
        return _parse_excel(filename, binary, parser_config, callback)
    if file_extension in TEXT_FILE_EXTENSIONS:
        return _parse_text(filename, binary, parser_config, callback)
    if file_extension in {".md", ".markdown"}:
        return _parse_markdown(filename, binary, parser_config, callback)
    if file_extension in {".htm", ".html"}:
        return _parse_html(filename, binary, callback)
    if file_extension == ".json":
        return _parse_json(filename, binary, parser_config, callback)
    if file_extension == ".doc":
        return _parse_legacy_doc(filename, binary, callback)

    raise NotImplementedError(
        f"Unsupported file type: {file_extension or '(no extension)'}"
    )


def _merge_and_tokenize(
    parsed_document,
    document_metadata,
    is_english,
    parser_config,
    section_only,
    filename,
):
    """把连续段落合并成 chunk，并生成 Elasticsearch 需要的分词字段。"""
    chunk_token_num = int(parser_config.get("chunk_token_num", 128))
    delimiter = parser_config.get("delimiter", DEFAULT_PARSER_CONFIG["delimiter"])

    # 表格不和普通段落混在一起，每张表单独生成 chunk。
    result = tokenize_table(
        parsed_document.tables,
        document_metadata,
        is_english,
    )
    start_time = timer()

    if parsed_document.uses_docx_merge:
        # DOCX 的 sections 里还带有图片，因此使用专门的合并函数。
        chunks, images = naive_merge_docx(
            parsed_document.sections,
            chunk_token_num,
            delimiter,
        )
        if section_only:
            return chunks

        result.extend(
            tokenize_chunks_docx(
                chunks,
                document_metadata,
                is_english,
                images,
            )
        )
    else:
        #从正文内容中拼接或切割出大小合适的 chunk，供后续检索使用
        chunks = naive_merge(
            parsed_document.sections,
            chunk_token_num,
            delimiter,
        )
        if section_only:
            return chunks

        #添加单个元素用append, 多个元素用extend
        result.extend(
            # 为正文chunk分词并裁剪图片
            tokenize_chunks(
                chunks,
                document_metadata,
                is_english,
                parsed_document.pdf_parser,
            )
        )

    logging.info("naive_merge(%s): %s", filename, timer() - start_time)
    return result


def chunk(filename, binary=None, from_page=0, to_page=100000,
          lang="Chinese", callback=None, **kwargs):
    """
    把一个文件解析成适合建立 RAG 索引的多个 chunk。

    可以把完整流程理解成三步：
    1. 根据扩展名选择解析器，把文件变成“段落和表格”；
    2. 按 token 数量把相邻段落合并成大小合适的 chunk；
    3. 为每个 chunk 生成分词字段，供 Elasticsearch 后续检索。

    :param filename: 文件路径，也会作为文档名称写入 chunk
    :param binary: 可选的文件二进制内容；没传时解析器会读取 filename
    :param from_page: PDF 开始页，页码从 0 开始
    :param to_page: PDF 结束页，不包含这一页
    :param lang: English 时使用英文分词方式，其余值按非英文处理
    :param callback: 可选的解析进度通知函数
    :param kwargs: 额外配置，目前使用 parser_config 和 section_only
    :return: 默认返回可写入 ES 的字典列表；section_only=True 时只返回文本块
    """
    # callback 表示解析进度通知函数
    progress_callback = callback or _empty_callback
    parser_config = _get_parser_config(kwargs)
    document_metadata = _build_document_metadata(filename)
    is_english = lang.lower() == "english"

    parsed_document = _parse_document(
        filename=filename,
        binary=binary,
        from_page=from_page,
        to_page=to_page,
        parser_config=parser_config,
        callback=progress_callback,
    )

    return _merge_and_tokenize(
        parsed_document=parsed_document,
        document_metadata=document_metadata,
        is_english=is_english,
        parser_config=parser_config,
        section_only=kwargs.get("section_only", False),
        filename=filename,
    )


if __name__ == "__main__":
    import sys


    def dummy(prog=None, msg=""):
        pass


    chunk(sys.argv[1], from_page=0, to_page=10, callback=dummy)
