import json
from types import SimpleNamespace
from unittest.mock import Mock

from PIL import Image
import pytest

from service.core.table_evidence import bind_table_rows
from service.core.figure_evidence import figure_evidence
from service.core.rag.nlp import tokenize_table
from service.core.rag.nlp import rag_tokenizer
from service.core import file_parse
from service.core.rag.utils import es_conn
from service.core.evidence_sufficiency import check_evidence_sufficiency
from service.core.retrieval import format_retrieved_chunk
from agent.tools.rag_search import _chunk_to_source
from agent.orchestrator import tool_results_to_retrieved_content
from agent.schemas import ToolResult


pytestmark = pytest.mark.unit


def test_multilevel_headers_rowspans_units_blank_and_zero():
    html = """<table><caption>表1 收入</caption>
    <tr><th colspan="4">单位：亿元</th></tr>
    <tr><th rowspan="2">业务</th><th rowspan="2">地区</th><th colspan="2">收入</th></tr>
    <tr><th>2023年</th><th>2024年</th></tr>
    <tr><td rowspan="2">发电</td><td>境内</td><td>0</td><td>-12.5</td></tr>
    <tr><td>境外</td><td></td><td>10</td></tr></table>"""
    rows = bind_table_rows(html)
    assert len(rows) == 2
    assert "收入 / 2023年：0" in rows[0]["content"]
    assert "收入 / 2024年：-12.5" in rows[0]["content"]
    assert "业务：发电；地区：境外" in rows[1]["content"]
    assert "收入 / 2023年：（空白）" in rows[1]["content"]
    assert all("单位：亿元" in row["content"] for row in rows)
    assert rows[1]["table_row_int"] == 5


def test_repeated_headers_and_row_header_do_not_swallow_data():
    rows = bind_table_rows("""<table><tr><th>指标</th><th>2023</th></tr>
    <tr><th scope="row">利润</th><td>1</td></tr>
    <tr><th>指标</th><th>2024</th></tr><tr><td>利润</td><td>2</td></tr></table>""")
    assert len(rows) == 2
    assert "2023：1" in rows[0]["content"]
    assert "2024：2" in rows[1]["content"]
    assert "2023" not in rows[1]["content"]


@pytest.mark.parametrize("html", [
    '<table><tr><td>收入</td><td>100</td></tr><tr><td>利润</td><td>20</td></tr></table>',
    '<table><tr><th colspan="oops">A</th></tr></table>',
    '<table><tr><td><table><tr><td>x</td></tr></table></td></tr></table>',
    '<table><tr><th rowspan="9">A</th></tr></table>',
])
def test_ambiguous_or_invalid_tables_are_preserved_without_inventing_headers(html):
    rows = bind_table_rows(html)
    assert rows[0]["content"].endswith(html)
    assert rows[0]["table_binding_kwd"] == "unbound"


def test_table_with_year_headers_and_equal_numbers_keeps_both_columns():
    rows = bind_table_rows('<table><tr><td>指标</td><td>2023</td><td>2024</td></tr>'
                           '<tr><td>A &amp; B</td><td>0</td><td>0</td></tr></table>')
    assert "2023：0；2024：0" in rows[0]["content"]
    assert "<td>A &amp; B</td>" in rows[0]["content"]


def test_key_value_table_with_ocr_header_tags_is_not_misbound():
    html = '<table><tr><th>信用评级</th><th>AAA</th></tr><tr><td>发行金额</td><td>100</td></tr></table>'
    assert bind_table_rows(html)[0]["table_binding_kwd"] == "unbound"
    assert bind_table_rows(html)[0]["content"].endswith(html)


def test_negative_data_mislabeled_as_th_does_not_replace_year_headers():
    html = '<table><tr><th>利润表（万元）</th><th>2024</th><th>2025E</th></tr>'
    html += '<tr><th>减值</th><th>(120)</th><th>(130)</th></tr>'
    html += '<tr><td>利润</td><td>200</td><td>300</td></tr></table>'
    rows = bind_table_rows(html)
    assert len(rows) == 2
    assert "2024：(120)；2025E：(130)" in rows[0]["content"]
    assert "2024：200；2025E：300" in rows[1]["content"]


def test_missing_group_header_spans_preserve_original_table():
    html = '<table><tr><th>指标</th><th>收入</th><th></th></tr>'
    html += '<tr><th></th><th>2024</th><th>2025</th></tr>'
    html += '<tr><td>产品</td><td>100</td><td>200</td></tr></table>'
    assert bind_table_rows(html)[0]["table_binding_kwd"] == "unbound"


def test_growth_row_carries_previous_metric_as_context():
    html = '<table><tr><th>指标</th><th>2025</th></tr><tr><td>收入</td><td>100</td></tr>'
    html += '<tr><td>(+/-%)</td><td>10%</td></tr></table>'
    assert "上一数据行指标：收入" in bind_table_rows(html)[1]["content"]


def test_body_section_labels_do_not_leak_to_following_sections():
    html = '<table><tr><th>指标</th><th>2024</th></tr><tr><td colspan="2">境内</td></tr>'
    html += '<tr><td>收入</td><td>10</td></tr><tr><td colspan="2">境外</td></tr>'
    html += '<tr><td>收入</td><td>20</td></tr></table>'
    rows = bind_table_rows(html)
    assert "境内" in rows[0]["content"]
    assert "境外" in rows[1]["content"]
    assert "境内" not in rows[1]["content"]


def test_docx_only_merges_actual_adjacent_merged_cells():
    from docx import Document
    from service.core.rag.app.naive import Docx
    doc = Document()
    table = doc.add_table(rows=1, cols=4)
    for cell in table.rows[0].cells:
        cell.text = "0"
    html = Docx._row_to_html(table.rows[0])
    assert html.count("<td>0</td>") == 4
    table.cell(0, 0).merge(table.cell(0, 1)).text = "A & B"
    html = Docx._row_to_html(table.rows[0])
    assert "colspan='2'" in html
    assert "A &amp; B" in html
    assert html.count("<td>0</td>") == 2


def test_disabled_vision_never_calls_provider(monkeypatch):
    monkeypatch.delenv("FIGURE_VISION_ENABLED", raising=False)
    client = Mock()
    result = figure_evidence(Image.new("RGB", (10, 10)), "图1 收入", client=client)
    assert result["visual_status_kwd"] == "disabled"
    assert "未提取视觉事实" in result["content"]
    client.chat.completions.create.assert_not_called()


@pytest.mark.parametrize("facts,status", [(["2024年收入为10亿元。"], "extracted"), ([], "empty"), ([123], "failed")])
def test_vision_sends_image_and_validates_facts(monkeypatch, facts, status):
    monkeypatch.setenv("FIGURE_VISION_ENABLED", "true")
    monkeypatch.setenv("FIGURE_VISION_MODEL", "test-vision")
    client = Mock()
    client.chat.completions.create.return_value = SimpleNamespace(choices=[
        SimpleNamespace(message=SimpleNamespace(content=json.dumps({"facts": facts})))])
    result = figure_evidence(Image.new("RGB", (10, 10)), "图1", client=client)
    assert result["visual_status_kwd"] == status
    call = client.chat.completions.create.call_args.kwargs
    assert call["model"] == "test-vision"
    assert call["messages"][1]["content"][1]["image_url"]["url"].startswith("data:image/png;base64,")
    if status != "extracted":
        assert "未提取视觉事实" in result["content"]


def test_vision_failure_does_not_abort_document(monkeypatch):
    monkeypatch.setenv("FIGURE_VISION_ENABLED", "true")
    monkeypatch.setenv("FIGURE_VISION_MODEL", "test-vision")
    client = Mock()
    client.chat.completions.create.side_effect = TimeoutError()
    assert figure_evidence(Image.new("RGB", (10, 10)), "图1", client=client)["visual_status_kwd"] == "failed"


def test_tokenization_retains_positions_and_figure_type(monkeypatch):
    monkeypatch.setenv("FIGURE_VISION_ENABLED", "false")
    monkeypatch.setattr(rag_tokenizer, "tokenize", lambda text: text)
    monkeypatch.setattr(rag_tokenizer, "fine_grained_tokenize", lambda text: text)
    pos = [(1, 10, 80, 20, 100)]
    image = Image.new("RGB", (10, 10))
    records = tokenize_table([((image, {"kind": "figure", "text": "图2 利润"}), pos)], {}, False)
    assert records[0]["page_num_int"] == [2]
    assert records[0]["position_int"] == [(2, 10, 80, 20, 100)]
    assert records[0]["evidence_type_kwd"] == "figure"
    assert "未提取视觉事实" in records[0]["content_with_weight"]


def test_chunk_ids_include_document_and_position_and_images_are_saved(monkeypatch, tmp_path):
    monkeypatch.setattr(file_parse, "batch_generate_embeddings", lambda texts: [[0.1] * 512 for _ in texts])
    monkeypatch.setattr(file_parse, "get_project_base_directory", lambda: str(tmp_path))
    item = {"content_with_weight": "相同内容", "content_ltks": "相同 内容", "content_sm_ltks": "相同 内容",
            "docnm_kwd": "/tmp/report.pdf", "title_tks": "report", "page_num_int": [2],
            "position_int": [(2, 1, 10, 5, 20)], "evidence_type_kwd": "table", "image": Image.new("RGB", (2, 2))}
    first = file_parse.process_items([item, item], "a.pdf", "1")
    second = file_parse.process_items([item], "b.pdf", "1")
    assert len({first[0]["id"], first[1]["id"], second[0]["id"]}) == 3
    assert first[0]["id"] == file_parse.process_items([item], "a.pdf", "1")[0]["id"]
    assert first[0]["page_num_int"] == [2]
    assert first[0]["docnm_kwd"] == "a.pdf"
    assert (tmp_path / first[0]["image_ref_kwd"]).is_file()
    json.dumps(first)


def test_reindex_writes_before_deleting_only_superseded_ids(monkeypatch):
    connection = es_conn.ESConnection()
    fake = Mock()
    fake.indices.exists.return_value = True
    fake.bulk.return_value = {"errors": False}
    monkeypatch.setattr(connection, "es", fake)
    monkeypatch.setattr(es_conn, "scan", lambda *a, **k: [{"_id": "old"}, {"_id": "retained"}])
    def insert(*args):
        fake.bulk.assert_not_called()
        return []
    monkeypatch.setattr(connection, "insert", insert)
    result = connection.replace_document([{"id": "retained", "doc_id": "d", "kb_id": "1"}], "1", "d")
    assert result == {"indexed": 1, "removed": 1}
    assert fake.bulk.call_args.kwargs["operations"] == [{"delete": {"_index": "1", "_id": "old"}}]


def test_failed_reindex_keeps_old_chunks(monkeypatch):
    connection = es_conn.ESConnection()
    fake = Mock()
    monkeypatch.setattr(connection, "es", fake)
    monkeypatch.setattr(es_conn, "scan", lambda *a, **k: [{"_id": "old"}])
    monkeypatch.setattr(connection, "insert", lambda *a: ["write failed"])
    with pytest.raises(RuntimeError, match="old chunks retained"):
        connection.replace_document([{"id": "new", "doc_id": "d", "kb_id": "1"}], "1", "d")
    fake.bulk.assert_not_called()


def test_bulk_non_timeout_exception_is_not_reported_as_success(monkeypatch):
    connection = es_conn.ESConnection()
    fake = Mock()
    fake.bulk.side_effect = RuntimeError("write denied")
    monkeypatch.setattr(connection, "es", fake)
    assert connection.insert([{"id": "id"}], "1") == ["write denied"]


@pytest.mark.parametrize("enabled", ["true", "false"])
def test_ocr_only_figure_cannot_answer_visual_question_even_if_verifier_disabled(monkeypatch, enabled):
    monkeypatch.setenv("RAG_EVIDENCE_SUFFICIENCY_ENABLED", enabled)
    decision = check_evidence_sufficiency("图3中哪一年为负？", [{
        "chunk_id": "fig", "content_with_weight": "图3 利润 2021 2022",
        "evidence_type_kwd": "figure", "visual_status_kwd": "disabled",
    }], client=Mock())
    assert decision.sufficient is False
    assert decision.source == "rule"


def test_retrieval_and_tool_source_keep_evidence_metadata():
    formatted = format_retrieved_chunk({"chunk_id": "fig", "content_with_weight": "图1",
        "evidence_type_kwd": "figure", "visual_status_kwd": "disabled", "page_num_int": [2]}, 1)
    source = _chunk_to_source(formatted, 1)
    assert source.metadata["visual_status_kwd"] == "disabled"
    assert source.metadata["page_num_int"] == [2]
    references = tool_results_to_retrieved_content([
        ToolResult(tool_name="search_knowledge_base", query="图1", content="图1", sources=[source])
    ])
    assert references[0]["visual_status_kwd"] == "disabled"
    assert references[0]["page_num_int"] == [2]
